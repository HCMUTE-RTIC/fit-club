"""API integration tests: isolated DB, real authorization and signed QR paths."""
import os
from io import BytesIO

os.environ['JWT_SECRET'] = 'test-only-secret-that-is-at-least-32-characters'
os.environ['DATABASE_URL'] = 'sqlite://'
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
import pytest
from sqlalchemy import create_engine
from sqlalchemy.pool import StaticPool
from sqlalchemy.orm import Session
from app import main as m


@pytest.fixture
def api(monkeypatch):
    engine = create_engine('sqlite://', connect_args={'check_same_thread': False}, poolclass=StaticPool)
    m.Base.metadata.create_all(engine)
    def database():
        with Session(engine) as session:
            yield session
    m.app.dependency_overrides[m.db] = database
    monkeypatch.setattr(m, 'engine', engine)
    monkeypatch.setattr(m.settings, 'superadmin_emails', '10000001@student.hcmute.edu.vn')
    with TestClient(m.app) as client:
        yield client
    m.app.dependency_overrides.clear()
    engine.dispose()


def auth(api, monkeypatch, sid='10000001', **extra):
    claims = {'sub': 'google-' + sid, 'email': sid + '@student.hcmute.edu.vn', 'hd': m.DOMAIN, 'email_verified': True, 'name': 'Student ' + sid, **extra}
    monkeypatch.setattr(m, 'verify_google', lambda credential: claims)
    response = api.post('/auth/google', json={'credential': 'test-google-credential'})
    assert response.status_code == 200, response.text
    return {'Authorization': 'Bearer ' + response.json()['access_token']}


def event(api, headers, title='Club Day'):
    r = api.post('/events', headers=headers, json={'title': title, 'starts_at': '2026-10-17T08:00:00+07:00', 'checkin_open': True})
    assert r.status_code == 201, r.text
    return r.json()['id']


def excel(rows):
    book = Workbook(); sheet = book.active
    for row in [('student_id', 'full_name', 'email'), *rows]:
        sheet.append(row)
    stream = BytesIO(); book.save(stream)
    return {'file': ('students.xlsx', stream.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}


@pytest.mark.parametrize('claims', [
    {'email_verified': False}, {'hd': 'gmail.com'}, {'email': '10000001@gmail.com'},
    {'email': 'abc@student.hcmute.edu.vn'}, {'sub': ''}, {'name': ''},
])
def test_google_identity_restrictions(api, monkeypatch, claims):
    payload = {'sub': 'google-id', 'email': '10000001@student.hcmute.edu.vn', 'email_verified': True, 'hd': m.DOMAIN, 'name': 'Name', **claims}
    monkeypatch.setattr(m, 'verify_google', lambda credential: payload)
    assert api.post('/auth/google', json={'credential': 'test-google-credential'}).status_code in (403, 422)
    assert api.get('/auth/me').status_code in (401, 403)


def setup_event(api, monkeypatch):
    root = auth(api, monkeypatch)
    first, second = event(api, root), event(api, root, 'Second')
    assert api.post(f'/events/{first}/admins', headers=root, json={'email': '10000002@student.hcmute.edu.vn', 'name': 'Admin'}).status_code == 201
    admin = auth(api, monkeypatch, '10000002')
    student = auth(api, monkeypatch, '10000003')
    assert api.post(f'/events/{first}/students', headers=admin, json={'email': '10000003@student.hcmute.edu.vn', 'name': 'Student'}).status_code == 201
    return root, admin, student, first, second


def test_roles_are_scoped_and_revocation_is_immediate(api, monkeypatch):
    root, admin, student, first, second = setup_event(api, monkeypatch)
    assert api.get('/auth/me', headers=admin).json()['role'] == 'admin'
    assert api.get('/auth/me', headers=student).json()['student_id'] == '10000003'
    assert [x['id'] for x in api.get('/events', headers=student).json()] == [first]
    for path in ('students', 'students/template', 'attendance/export'):
        assert api.get(f'/events/{second}/{path}', headers=admin).status_code == 403
        assert api.get(f'/events/{first}/{path}', headers=student).status_code == 403
    assert api.post('/events', headers=admin, json={'title': 'Forbidden', 'starts_at': '2026-10-17T08:00:00+07:00'}).status_code == 403
    assert api.post(f'/events/{first}/admins', headers=admin, json={'email': '10000004@student.hcmute.edu.vn', 'name': 'No'}).status_code == 403
    aid = api.get('/auth/me', headers=admin).json()['id']
    assert api.delete(f'/events/{first}/admins/{aid}', headers=root).status_code == 204
    assert api.get(f'/events/{first}/students', headers=admin).status_code == 403


def test_excel_template_duplicates_atomic_errors_and_search(api, monkeypatch):
    _, admin, _, first, _ = setup_event(api, monkeypatch)
    url = f'/events/{first}/students'
    template = api.get(url + '/template', headers=admin)
    assert list(load_workbook(BytesIO(template.content)).active.values) == [('student_id', 'full_name', 'email')]
    rows = [('10000004', 'Four', ''), ('10000004', 'Four again', ''), ('10000003', 'Existing', '')]
    response = api.post(url + '/import', headers=admin, files=excel(rows)).json()
    assert response['imported'] == 1 and len(response['duplicates']) == 2
    assert api.post(url + '/import', headers=admin, files=excel(rows)).json()['imported'] == 0
    invalid = [('10000005', 'Five', ''), ('10000006', 'Mismatch', '10000007@student.hcmute.edu.vn')]
    response = api.post(url + '/import', headers=admin, files=excel(invalid)).json()
    assert response['imported'] == 0 and response['errors']
    assert api.get(url, headers=admin, params={'search': '10000005'}).json() == []
    assert len(api.get(url, headers=admin, params={'search': '10000004'}).json()) == 1
    assert api.post(url + '/import', headers=admin, files={'file': ('bad.xlsx', b'not excel')}).status_code == 422


def test_qr_security_repeat_scan_and_export(api, monkeypatch):
    root, admin, student, first, second = setup_event(api, monkeypatch)
    assert api.get(f'/events/{second}/qr', headers=student).status_code == 403
    qr = api.get(f'/events/{first}/qr', headers=student).json()
    assert qr['image_data_url'].startswith('data:image/png;base64,')
    url = f'/events/{first}/attendance/scan'
    assert api.post(url, headers=student, json={'token': qr['token']}).status_code == 403
    assert api.post(f'/events/{second}/attendance/scan', headers=root, json={'token': qr['token']}).status_code == 403
    assert api.get('/auth/me', headers={'Authorization': 'Bearer ' + qr['token']}).status_code == 401
    assert api.post(url, headers=admin, json={'token': student['Authorization'][7:]}).status_code == 401
    sid = api.get('/auth/me', headers=student).json()['id']
    expired = m.token({'sub': str(sid), 'kind': 'checkin', 'event_id': first}, -1)
    assert api.post(url, headers=admin, json={'token': expired}).status_code == 401
    for repeat in (False, True):
        response = api.post(url, headers=admin, json={'token': qr['token']})
        assert response.status_code == 200, response.text
        assert response.json()['already_checked_in'] is repeat
    records = api.get(f'/events/{first}/students', headers=admin, params={'attended': True}).json()
    assert len(records) == 1 and records[0]['checked_by']
    exported = api.get(f'/events/{first}/attendance/export', headers=admin)
    rows = list(load_workbook(BytesIO(exported.content)).active.values)
    assert rows[1][0] == '10000003' and rows[1][3] == 'Có mặt'
    assert rows[1][5] == '10000002@student.hcmute.edu.vn'


def test_export_never_executes_names_as_excel_formulas(api, monkeypatch):
    root = auth(api, monkeypatch)
    eid = event(api, root)
    assert api.post(f'/events/{eid}/students', headers=root, json={'email': '10000008@student.hcmute.edu.vn', 'name': '=HYPERLINK("https://example.com")'}).status_code == 201
    result = api.get(f'/events/{eid}/attendance/export', headers=root)
    assert load_workbook(BytesIO(result.content)).active['B2'].data_type == 's'


def test_closed_checkin_and_unregistered_signed_qr(api, monkeypatch):
    root, admin, student, first, _ = setup_event(api, monkeypatch)
    uid = api.get('/auth/me', headers=student).json()['id']
    unknown = m.token({'sub': str(uid + 999), 'kind': 'checkin', 'event_id': first}, 120)
    assert api.post(f'/events/{first}/attendance/scan', headers=admin, json={'token': unknown}).status_code == 403
    qr = api.get(f'/events/{first}/qr', headers=student).json()['token']
    update = api.put(f'/events/{first}', headers=root, json={'title': 'Closed', 'starts_at': '2026-10-17T08:00:00+07:00', 'checkin_open': False})
    assert update.status_code == 200
    assert api.get(f'/events/{first}/qr', headers=student).status_code == 409
    assert api.post(f'/events/{first}/attendance/scan', headers=admin, json={'token': qr}).status_code == 409


def test_google_existing_identity_cannot_be_rebound(api, monkeypatch):
    auth(api, monkeypatch, '10000003')
    monkeypatch.setattr(m, 'verify_google', lambda credential: {'sub': 'different-identity', 'email': '10000003@student.hcmute.edu.vn', 'hd': m.DOMAIN, 'email_verified': True, 'name': 'Other'})
    assert api.post('/auth/google', json={'credential': 'test-google-credential'}).status_code == 403


def test_malformed_workbook_xml_is_validation_error(api, monkeypatch):
    import zipfile
    root = auth(api, monkeypatch)
    eid = event(api, root)
    valid = excel([])['file'][1]
    output = BytesIO()
    with zipfile.ZipFile(BytesIO(valid)) as original, zipfile.ZipFile(output, 'w') as modified:
        for name in original.namelist():
            modified.writestr(name, b'<broken' if name == 'xl/workbook.xml' else original.read(name))
    response = api.post(f'/events/{eid}/students/import', headers=root, files={'file': ('bad.xlsx', output.getvalue())})
    assert response.status_code == 422, response.text


def test_cors_preflight_supports_event_update(api):
    response = api.options('/events/1', headers={'Origin': m.settings.cors_origins.split(',')[0].strip(), 'Access-Control-Request-Method': 'PUT', 'Access-Control-Request-Headers': 'authorization,content-type'})
    assert response.status_code == 200
    assert 'PUT' in response.headers['access-control-allow-methods']


def test_grant_superadmin_before_login_and_immediate_revoke(api, monkeypatch):
    root = auth(api, monkeypatch)
    granted = api.post('/superadmins', headers=root, json={'email': '10000009@student.hcmute.edu.vn', 'name': 'New Superadmin'})
    assert granted.status_code == 201, granted.text
    uid = granted.json()['id']
    new_root = auth(api, monkeypatch, '10000009')
    assert api.get('/auth/me', headers=new_root).json()['role'] == 'superadmin'
    event(api, new_root, 'Created by newly granted superadmin')
    assert uid in [p['id'] for p in api.get('/superadmins', headers=root).json()['users']]
    assert api.post('/superadmins', headers=root, json={'email': '10000009@student.hcmute.edu.vn', 'name': 'Duplicate'}).status_code == 409
    assert api.delete(f'/superadmins/{uid}', headers=root).status_code == 204
    assert api.get('/auth/me', headers=new_root).json()['role'] == 'student'
    assert api.post('/events', headers=new_root, json={'title': 'Denied after revoke', 'starts_at': '2026-10-17T08:00:00+07:00'}).status_code == 403
    assert api.get('/superadmins', headers=new_root).status_code == 403
    bootstrap_id = api.get('/auth/me', headers=root).json()['id']
    assert api.delete(f'/superadmins/{bootstrap_id}', headers=root).status_code == 409


def test_students_and_event_admins_cannot_manage_superadmins(api, monkeypatch):
    _, admin, student, _, _ = setup_event(api, monkeypatch)
    for headers in (admin, student):
        assert api.get('/superadmins', headers=headers).status_code == 403
        assert api.post('/superadmins', headers=headers, json={'email': '10000009@student.hcmute.edu.vn', 'name': 'Denied'}).status_code == 403
        assert api.delete('/superadmins/1', headers=headers).status_code == 403
