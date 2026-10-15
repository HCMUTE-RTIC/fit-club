"""Independent event-admin actor acceptance checks; isolated in-memory DB."""
from io import BytesIO
import pytest
from openpyxl import load_workbook
from test_api import api, auth, event, excel, setup_event

@pytest.mark.parametrize('suffix', ['students', 'students/template', 'attendance/export'])
def test_admin_cannot_read_other_event(api, monkeypatch, suffix):
    _, admin, _, _, other = setup_event(api, monkeypatch)
    response = api.get(f'/events/{other}/{suffix}', headers=admin)
    assert response.status_code == 403, response.text


def test_admin_cannot_mutate_other_event(api, monkeypatch):
    _, admin, student, own, other = setup_event(api, monkeypatch)
    qr = api.get(f'/events/{own}/qr', headers=student).json()['token']
    for suffix, kwargs in [
        ('students', {'json': {'email': '10000004@student.hcmute.edu.vn', 'name': 'Denied'}}),
        ('students/import', {'files': excel([('10000004', 'Denied', '')])}),
        ('attendance/scan', {'json': {'token': qr}}),
    ]:
        response = api.post(f'/events/{other}/{suffix}', headers=admin, **kwargs)
        assert response.status_code == 403, response.text


def test_admin_has_no_global_role_management(api, monkeypatch):
    _, admin, _, own, _ = setup_event(api, monkeypatch)
    person = {'email': '10000004@student.hcmute.edu.vn', 'name': 'Denied'}
    attempts = [
        ('post', '/events', {'json': {'title': 'No', 'starts_at': '2026-10-17T08:00:00+07:00'}}),
        ('put', f'/events/{own}', {'json': {'title': 'No', 'starts_at': '2026-10-17T08:00:00+07:00'}}),
        ('post', f'/events/{own}/admins', {'json': person}),
        ('get', f'/events/{own}/admins', {}),
        ('delete', f'/events/{own}/admins/1', {}),
        ('post', '/superadmins', {'json': person}),
        ('get', '/superadmins', {}),
        ('delete', '/superadmins/1', {}),
    ]
    for method, url, kwargs in attempts:
        response = getattr(api, method)(url, headers=admin, **kwargs)
        assert response.status_code == 403, (method, url, response.text)


def test_admin_template_roundtrip_duplicate_search_and_atomic_invalid(api, monkeypatch):
    _, admin, _, own, _ = setup_event(api, monkeypatch)
    url = f'/events/{own}/students'
    template = api.get(url + '/template', headers=admin)
    assert template.status_code == 200
    book = load_workbook(BytesIO(template.content))
    assert list(book.active.values) == [('student_id', 'full_name', 'email')]
    for row in [('10000004', 'Nguyễn Bốn', ''), ('10000004', 'Duplicate', ''), ('10000003', 'Existing', '')]:
        book.active.append(row)
    buffer = BytesIO(); book.save(buffer)
    result = api.post(url + '/import', headers=admin, files={'file': ('filled.xlsx', buffer.getvalue())})
    assert result.status_code == 200, result.text
    assert result.json()['imported'] == 1
    assert {item['reason'] for item in result.json()['duplicates']} == {'Trùng trong file', 'Đã có trong sự kiện'}
    invalid = api.post(url + '/import', headers=admin, files=excel([
        ('10000005', 'Valid but must rollback', ''), ('10000006', 'Mismatch', '10000007@student.hcmute.edu.vn')]))
    assert invalid.status_code == 200
    assert invalid.json()['imported'] == 0 and invalid.json()['errors'][0]['row'] == 3
    assert api.get(url, headers=admin, params={'search': '10000005'}).json() == []
    found = api.get(url, headers=admin, params={'search': '000004'}).json()
    assert len(found) == 1 and found[0]['student_id'] == '10000004'
    assert api.get(url, headers=admin, params={'search': '%'}).json() == []
    assert api.post(url + '/import', headers=admin, files=excel([('10000004', 'Repeated', '')])).json()['imported'] == 0


def test_admin_attendance_repeat_filter_export_and_close(api, monkeypatch):
    root, admin, student, own, other = setup_event(api, monkeypatch)
    students_url = f'/events/{own}/students'
    assert api.post(students_url, headers=admin, json={'email': '10000004@student.hcmute.edu.vn', 'name': 'Absent'}).status_code == 201
    qr = api.get(f'/events/{own}/qr', headers=student).json()['token']
    scan_url = f'/events/{own}/attendance/scan'
    first = api.post(scan_url, headers=admin, json={'token': qr})
    repeated = api.post(scan_url, headers=admin, json={'token': qr})
    assert first.status_code == repeated.status_code == 200
    assert first.json()['already_checked_in'] is False
    assert repeated.json()['already_checked_in'] is True
    assert first.json()['checked_at'] == repeated.json()['checked_at']
    assert first.json()['checked_by'] == api.get('/auth/me', headers=admin).json()['id']
    assert [row['student_id'] for row in api.get(students_url, headers=admin, params={'attended': True}).json()] == ['10000003']
    assert [row['student_id'] for row in api.get(students_url, headers=admin, params={'attended': False}).json()] == ['10000004']
    export = api.get(f'/events/{own}/attendance/export', headers=admin)
    assert export.status_code == 200
    rows = list(load_workbook(BytesIO(export.content)).active.values)
    assert len(rows) == 3 and rows[1][3] == 'Có mặt' and rows[2][3] == 'Chưa điểm danh'
    assert rows[1][5] == '10000002@student.hcmute.edu.vn'
    filtered = api.get(f'/events/{own}/attendance/export', headers=admin, params={'attended': True, 'search': '10000003'})
    assert len(list(load_workbook(BytesIO(filtered.content)).active.values)) == 2
    assert api.put(f'/events/{own}', headers=root, json={'title': 'Closed', 'starts_at': '2026-10-17T08:00:00+07:00', 'checkin_open': False}).status_code == 200
    assert api.post(scan_url, headers=admin, json={'token': qr}).status_code == 409


def test_admin_revocation_invalidates_existing_session_for_every_operation(api, monkeypatch):
    root, admin, student, own, _ = setup_event(api, monkeypatch)
    qr = api.get(f'/events/{own}/qr', headers=student).json()['token']
    uid = api.get('/auth/me', headers=admin).json()['id']
    assert api.delete(f'/events/{own}/admins/{uid}', headers=root).status_code == 204
    assert api.get('/auth/me', headers=admin).json()['role'] == 'student'
    assert api.get('/events', headers=admin).json() == []
    for suffix in ('students', 'students/template', 'attendance/export'):
        assert api.get(f'/events/{own}/{suffix}', headers=admin).status_code == 403
    assert api.post(f'/events/{own}/students', headers=admin, json={'email': '10000004@student.hcmute.edu.vn', 'name': 'Denied'}).status_code == 403
    assert api.post(f'/events/{own}/students/import', headers=admin, files=excel([])).status_code == 403
    assert api.post(f'/events/{own}/attendance/scan', headers=admin, json={'token': qr}).status_code == 403
