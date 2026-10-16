"""Independent superadmin actor tests, entirely isolated from the live DB."""
from datetime import datetime, timezone
from io import BytesIO
from openpyxl import load_workbook
from test_api import api, auth, event, excel
import pytest


def test_event_full_lifecycle_and_timezone(api, monkeypatch):
    root = auth(api, monkeypatch)
    payload = {'title': 'FIT 2026', 'description': 'Mô tả sự kiện', 'image_url': '/assets/banner.png', 'starts_at': '2026-10-17T08:00:00+07:00'}
    r = api.post('/events', headers=root, json=payload)
    assert r.status_code == 201
    eid = r.json()['id']
    assert datetime.fromisoformat(r.json()['starts_at']) == datetime(2026, 10, 17, 1, tzinfo=timezone.utc)
    payload.update(title='Updated FIT', image_url='https://example.com/event.png', checkin_open=True)
    r = api.put(f'/events/{eid}', headers=root, json=payload)
    assert r.status_code == 200
    assert r.json()['checkin_open'] is True
    assert api.get('/events', headers=root).json()[0] == r.json()
    assert api.put('/events/9999', headers=root, json=payload).status_code == 404


@pytest.mark.parametrize('change', [
    {'starts_at': '2026-10-17T08:00:00'}, {'title': ''},
    {'image_url': 'http://example.com/a.png'}, {'image_url': 'javascript:alert(1)'},
    {'image_url': '//evil.example/image.png'}, {'description': 'x' * 20001},
])
def test_invalid_events_do_not_persist(api, monkeypatch, change):
    root = auth(api, monkeypatch)
    payload = {'title': 'FIT', 'starts_at': '2026-10-17T08:00:00+07:00', **change}
    assert api.post('/events', headers=root, json=payload).status_code == 422
    assert api.get('/events', headers=root).json() == []


def test_multiple_admins_scoped_revoke_with_existing_token(api, monkeypatch):
    root = auth(api, monkeypatch)
    first, second = event(api, root), event(api, root, 'Other')
    ids = []
    for sid in ('20000001', '20000002'):
        body = {'email': sid + '@student.hcmute.edu.vn', 'name': sid}
        r = api.post(f'/events/{first}/admins', headers=root, json=body)
        assert r.status_code == 201
        ids.append(r.json()['id'])
        assert api.post(f'/events/{first}/admins', headers=root, json=body).status_code == 409
    assert len(api.get(f'/events/{first}/admins', headers=root).json()) == 2
    admin = auth(api, monkeypatch, '20000001')
    assert api.get(f'/events/{first}/students', headers=admin).status_code == 200
    assert api.get(f'/events/{second}/students', headers=admin).status_code == 403
    assert api.delete(f'/events/{first}/admins/{ids[0]}', headers=root).status_code == 204
    assert api.get(f'/events/{first}/students', headers=admin).status_code == 403
    assert api.delete(f'/events/{first}/admins/{ids[0]}', headers=root).status_code == 404
    assert len(api.get(f'/events/{first}/admins', headers=root).json()) == 1


def test_superadmin_grant_revoke_bootstrap_and_retained_event_admin(api, monkeypatch):
    root = auth(api, monkeypatch)
    eid = event(api, root)
    body = {'email': '20000003@student.hcmute.edu.vn', 'name': 'Second Super'}
    granted = api.post('/superadmins', headers=root, json=body)
    assert granted.status_code == 201
    uid = granted.json()['id']
    second = auth(api, monkeypatch, '20000003')
    assert api.get('/auth/me', headers=second).json()['role'] == 'superadmin'
    assert api.post(f'/events/{eid}/admins', headers=root, json=body).status_code == 201
    assert api.post('/superadmins', headers=root, json=body).status_code == 409
    assert len(api.get('/superadmins', headers=second).json()['users']) == 2
    assert api.delete(f'/superadmins/{uid}', headers=root).status_code == 204
    assert api.get('/auth/me', headers=second).json()['role'] == 'admin'
    assert api.get(f'/events/{eid}/students', headers=second).status_code == 200
    assert api.get('/superadmins', headers=second).status_code == 403
    bootstrap_id = api.get('/auth/me', headers=root).json()['id']
    assert api.delete(f'/superadmins/{bootstrap_id}', headers=root).status_code == 409
    assert api.delete('/superadmins/9999', headers=root).status_code == 404


def test_superadmin_bulk_template_duplicates_search_export(api, monkeypatch):
    root = auth(api, monkeypatch)
    eid = event(api, root)
    path = f'/events/{eid}/students'
    template = api.get(path + '/template', headers=root)
    assert template.status_code == 200
    assert tuple(load_workbook(BytesIO(template.content)).active.values) == (('student_id', 'full_name', 'email'),)
    rows = [('20000010', 'Nguyễn A', ''), ('20000011', 'Trần B', ''), ('20000010', 'Duplicate', '')]
    imported = api.post(path + '/import', headers=root, files=excel(rows))
    assert imported.status_code == 200
    assert imported.json()['imported'] == 2
    assert len(imported.json()['duplicates']) == 1
    assert len(api.get(path, headers=root, params={'search': '20000010'}).json()) == 1
    assert len(api.get(path, headers=root, params={'attended': False}).json()) == 2
    exported = api.get(f'/events/{eid}/attendance/export', headers=root)
    assert exported.status_code == 200
    assert len(list(load_workbook(BytesIO(exported.content)).active.values)) == 3
    repeated = api.post(path + '/import', headers=root, files=excel(rows)).json()
    assert repeated['imported'] == 0 and len(repeated['duplicates']) == 3


@pytest.mark.parametrize('endpoint,expected', [('/superadmins', 201), ('/events/1/admins', 201), ('/events/1/students', 422)])
def test_external_email_is_allowed_for_staff_but_not_student_rosters(api, monkeypatch, endpoint, expected):
    root = auth(api, monkeypatch)
    event(api, root)
    assert api.post(endpoint, headers=root, json={'email': 'person@gmail.com', 'name': 'Staff'}).status_code == expected


def test_blank_person_name_rejected(api, monkeypatch):
    root = auth(api, monkeypatch)
    eid = event(api, root)
    assert api.post(f'/events/{eid}/admins', headers=root, json={'email': '20000020@student.hcmute.edu.vn', 'name': '   '}).status_code == 422


def test_malformed_https_image_rejected(api, monkeypatch):
    root = auth(api, monkeypatch)
    assert api.post('/events', headers=root, json={'title': 'Bad image', 'image_url': 'https://', 'starts_at': '2026-10-17T08:00:00+07:00'}).status_code == 422
