from contextlib import asynccontextmanager
from datetime import datetime, timedelta, timezone
from io import BytesIO
import re
import hashlib
import secrets
import zipfile

import jwt
import qrcode
from fastapi import Depends, FastAPI, File, HTTPException, Query, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response
from fastapi.staticfiles import StaticFiles
from pathlib import Path
from xml.etree.ElementTree import ParseError
from fastapi.security import HTTPAuthorizationCredentials, HTTPBearer
from google.auth.transport.requests import Request
from google.oauth2 import id_token
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import BaseModel, Field, field_validator, HttpUrl, TypeAdapter
from pydantic_settings import BaseSettings, SettingsConfigDict
from sqlalchemy import Boolean, DateTime, ForeignKey, String, UniqueConstraint, create_engine, event as sql_event, or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import DeclarativeBase, Mapped, Session, mapped_column, object_session


class Settings(BaseSettings):
    model_config = SettingsConfigDict(env_file='.env', extra='ignore')
    database_url: str = 'sqlite:///./fitclub.db'
    jwt_secret: str = Field(min_length=32)
    google_client_id: str = ''
    student_email_only: bool = True
    superadmin_emails: str = ''
    cors_origins: str = 'http://localhost:3000,http://localhost:5173'
    access_minutes: int = 60
    qr_seconds: int = 120


settings = Settings()
DOMAIN = 'student.hcmute.edu.vn'
EMAIL_RE = re.compile(r'^([0-9]{6,15})@student\.hcmute\.edu\.vn$')
engine = create_engine(settings.database_url, connect_args={'check_same_thread': False} if settings.database_url.startswith('sqlite') else {})
if settings.database_url.startswith('sqlite'):
    @sql_event.listens_for(engine, 'connect')
    def sqlite_fk(connection, _):
        connection.execute('PRAGMA foreign_keys=ON')


class Base(DeclarativeBase):
    pass


class User(Base):
    __tablename__ = 'users'
    id: Mapped[int] = mapped_column(primary_key=True)
    email: Mapped[str] = mapped_column(String(254), unique=True)
    student_id: Mapped[str] = mapped_column(String(15), unique=True)
    name: Mapped[str] = mapped_column(String(200))
    google_sub: Mapped[str | None] = mapped_column(String(255), unique=True)


class SuperAdmin(Base):
    __tablename__ = 'superadmins'
    user_id: Mapped[int] = mapped_column(ForeignKey('users.id'), primary_key=True)


class Event(Base):
    __tablename__ = 'events'
    id: Mapped[int] = mapped_column(primary_key=True)
    title: Mapped[str] = mapped_column(String(200))
    description: Mapped[str]
    image_url: Mapped[str | None]
    starts_at: Mapped[datetime] = mapped_column(DateTime(timezone=True))
    checkin_open: Mapped[bool] = mapped_column(Boolean, default=False)


class EventAdmin(Base):
    __tablename__ = 'event_admins'
    event_id: Mapped[int] = mapped_column(ForeignKey('events.id'), primary_key=True)
    user_id: Mapped[int] = mapped_column(ForeignKey('users.id'), primary_key=True)


class Registration(Base):
    __tablename__ = 'registrations'
    event_id: Mapped[int] = mapped_column(ForeignKey('events.id'), primary_key=True)
    user_id: Mapped[int] = mapped_column(ForeignKey('users.id'), primary_key=True)
    registered_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))


class Attendance(Base):
    __tablename__ = 'attendance'
    __table_args__ = (UniqueConstraint('event_id', 'user_id'),)
    id: Mapped[int] = mapped_column(primary_key=True)
    event_id: Mapped[int] = mapped_column(ForeignKey('events.id'))
    user_id: Mapped[int] = mapped_column(ForeignKey('users.id'))
    checked_by: Mapped[int] = mapped_column(ForeignKey('users.id'))
    checked_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))


@asynccontextmanager
async def lifespan(app):
    Base.metadata.create_all(engine)
    yield


app = FastAPI(title='FIT Club Day API', version='1.0.0', lifespan=lifespan)
app.mount('/assets', StaticFiles(directory=Path(__file__).resolve().parents[2] / 'assets'), name='assets')
_cors_list = [x.strip() for x in settings.cors_origins.split(',') if x.strip()]
app.add_middleware(
    CORSMiddleware,
    allow_origins=['*'] if '*' in _cors_list or not _cors_list else _cors_list,
    allow_credentials=True,
    allow_methods=['*'],
    allow_headers=['*'],
)
bearer = HTTPBearer()


def db():
    with Session(engine) as session:
        yield session


def student_email(value):
    if not isinstance(value, str):
        raise HTTPException(422, 'Email không hợp lệ')
    value = value.strip().lower()
    match = EMAIL_RE.fullmatch(value)
    if match:
        return value, match.group(1)
    if settings.student_email_only:
        raise HTTPException(422, 'Email phải có dạng <MSSV số>@student.hcmute.edu.vn')
    if len(value) > 254 or not re.fullmatch(r"[a-z0-9.!#$%&'*+/=?^_`{|}~-]+@[a-z0-9](?:[a-z0-9-]*[a-z0-9])?(?:\.[a-z0-9](?:[a-z0-9-]*[a-z0-9])?)+", value):
        raise HTTPException(422, 'Email không hợp lệ')
    # External accounts have a stable identifier, never a fabricated numeric MSSV.
    return value, 'g' + hashlib.sha256(value.encode()).hexdigest()[:14]


def is_super(user):
    if user.email in {x.strip().lower() for x in settings.superadmin_emails.split(',') if x.strip()}:
        return True
    attached = object_session(user)
    if attached is not None:
        return attached.get(SuperAdmin, user.id) is not None
    with Session(engine) as session:
        return session.get(SuperAdmin, user.id) is not None


def token(payload, seconds):
    now = datetime.now(timezone.utc)
    return jwt.encode({**payload, 'iat': now, 'exp': now + timedelta(seconds=seconds), 'iss': 'fitclub', 'jti': secrets.token_urlsafe(16)}, settings.jwt_secret, algorithm='HS256')


def decode(value, kind):
    try:
        claims = jwt.decode(value, settings.jwt_secret, algorithms=['HS256'], issuer='fitclub', options={'require': ['exp', 'iat', 'sub', 'kind', 'jti']})
        if claims['kind'] != kind:
            raise ValueError()
        return claims
    except (jwt.PyJWTError, ValueError):
        raise HTTPException(401, 'Token không hợp lệ hoặc đã hết hạn')


def current(credentials: HTTPAuthorizationCredentials = Depends(bearer), session: Session = Depends(db)):
    claims = decode(credentials.credentials, 'access')
    try:
        user = session.get(User, int(claims['sub']))
    except (ValueError, TypeError):
        user = None
    if not user:
        raise HTTPException(401, 'Không tìm thấy tài khoản')
    return user


def superuser(user: User = Depends(current)):
    if not is_super(user):
        raise HTTPException(403, 'Chỉ superadmin được phép')
    return user


def get_event(session, event_id):
    item = session.get(Event, event_id)
    if not item:
        raise HTTPException(404, 'Không tìm thấy sự kiện')
    return item


def manage(session, user, event_id):
    item = get_event(session, event_id)
    if not is_super(user) and session.scalar(select(Club.id).where(Club.event_id == event_id).limit(1)) is not None:
        raise HTTPException(403, 'Sự kiện dùng phân quyền CLB. Hãy quản lý qua gian hàng được gán.')
    if not is_super(user) and not session.get(EventAdmin, (event_id, user.id)):
        raise HTTPException(403, 'Bạn không phải admin của sự kiện này')
    return item


def profile(user):
    return {'id': user.id, 'email': user.email, 'student_id': user.student_id, 'name': user.name}


def event_data(item):
    start = item.starts_at
    if start.tzinfo is None:
        start = start.replace(tzinfo=timezone.utc)
    return {'id': item.id, 'title': item.title, 'description': item.description, 'image_url': item.image_url, 'starts_at': start, 'checkin_open': item.checkin_open}


class GoogleLogin(BaseModel):
    credential: str = Field(min_length=10, max_length=10000)


def verify_google(credential):
    if not settings.google_client_id:
        raise HTTPException(503, 'Chưa cấu hình GOOGLE_CLIENT_ID')
    try:
        return id_token.verify_oauth2_token(credential, Request(), settings.google_client_id)
    except Exception as exc:
        # Google certificate/network errors also fail closed.
        raise HTTPException(401, 'Không xác minh được Google ID token') from exc


@app.post('/auth/google')
def login(body: GoogleLogin, session: Session = Depends(db)):
    claims = verify_google(body.credential)
    if claims.get('email_verified') is not True or (settings.student_email_only and claims.get('hd') != DOMAIN) or not claims.get('sub'):
        raise HTTPException(403, 'Chỉ tài khoản Google sinh viên HCMUTE đã xác minh được phép')
    email, sid = student_email(claims.get('email', ''))
    user = session.scalar(select(User).where(User.google_sub == claims['sub']))
    if user and user.email != email:
        raise HTTPException(403, 'Danh tính Google đã đổi email; cần liên hệ quản trị')
    user = user or session.scalar(select(User).where(User.email == email))
    if user and user.google_sub and user.google_sub != claims['sub']:
        raise HTTPException(403, 'Email đã gắn với danh tính Google khác')
    name = str(claims.get('name') or '').strip()[:200]
    if not name:
        raise HTTPException(422, 'Google không cung cấp họ tên')
    if not user:
        user = User(email=email, student_id=sid, name=name)
        session.add(user)
    user.google_sub = claims['sub']
    user.name = name
    session.commit()
    return {'access_token': token({'sub': str(user.id), 'kind': 'access'}, settings.access_minutes * 60), 'token_type': 'bearer', 'expires_in': settings.access_minutes * 60}


@app.get('/auth/me')
def me(user: User = Depends(current), session: Session = Depends(db)):
    ids = list(session.scalars(select(EventAdmin.event_id).where(EventAdmin.user_id == user.id)))
    club_ids = list(session.scalars(select(ClubAdmin.club_id).where(ClubAdmin.user_id == user.id)))
    return {**profile(user), 'role': 'superadmin' if is_super(user) else 'admin' if ids or club_ids else 'student', 'admin_event_ids': ids, 'admin_club_ids': club_ids}


class EventInput(BaseModel):
    @field_validator('title')
    @classmethod
    def clean_title(cls, value):
        if not value.strip():
            raise ValueError('Tiêu đề không được trống')
        return value.strip()

    title: str = Field(min_length=1, max_length=200)
    description: str = Field(default='', max_length=20000)
    image_url: str | None = Field(default=None, max_length=2000)
    starts_at: datetime
    checkin_open: bool = False

    @field_validator('starts_at')
    @classmethod
    def aware(cls, value):
        if value.tzinfo is None:
            raise ValueError('Thời gian phải có múi giờ, ví dụ +07:00')
        return value.astimezone(timezone.utc)

    @field_validator('image_url')
    @classmethod
    def image(cls, value):
        if value is None:
            return None
        if value.startswith('/assets/') and len(value) > len('/assets/') and not any(c in value for c in ('\\', '?', '#')) and '..' not in value.split('/'):
            return value
        parsed = TypeAdapter(HttpUrl).validate_python(value)
        if parsed.scheme != 'https' or parsed.username or parsed.password:
            raise ValueError('Ảnh phải là URL https hợp lệ hoặc đường dẫn /assets/')
        return str(parsed)


@app.post('/events', status_code=201)
def create_event(body: EventInput, user: User = Depends(superuser), session: Session = Depends(db)):
    item = Event(**body.model_dump())
    session.add(item)
    session.commit()
    return event_data(item)


@app.put('/events/{event_id}')
def update_event(event_id: int, body: EventInput, user: User = Depends(superuser), session: Session = Depends(db)):
    item = get_event(session, event_id)
    for key, value in body.model_dump().items():
        setattr(item, key, value)
    session.commit()
    return event_data(item)


@app.get('/events')
def events(user: User = Depends(current), session: Session = Depends(db)):
    query = select(Event).order_by(Event.starts_at.desc())
    if not is_super(user):
        query = query.where(or_(Event.id.in_(select(Registration.event_id).where(Registration.user_id == user.id)), Event.id.in_(select(EventAdmin.event_id).where(EventAdmin.user_id == user.id)), Event.id.in_(select(Club.event_id).join(ClubAdmin, ClubAdmin.club_id == Club.id).where(ClubAdmin.user_id == user.id))))
    return [event_data(e) for e in session.scalars(query)]


class Person(BaseModel):
    email: str
    name: str = Field(min_length=1, max_length=200)

    @field_validator('name')
    @classmethod
    def clean_name(cls, value):
        if not value.strip():
            raise ValueError('Họ tên không được trống')
        return value.strip()


def ensure_user(session, email, name):
    email, sid = student_email(email)
    user = session.scalar(select(User).where(User.email == email))
    if not user:
        user = User(email=email, student_id=sid, name=name.strip())
        session.add(user)
        session.flush()
    return user


@app.post('/events/{event_id}/admins', status_code=201)
def add_admin(event_id: int, body: Person, user: User = Depends(superuser), session: Session = Depends(db)):
    get_event(session, event_id)
    person = ensure_user(session, body.email, body.name)
    if session.get(EventAdmin, (event_id, person.id)):
        raise HTTPException(409, 'Admin đã được gán')
    session.add(EventAdmin(event_id=event_id, user_id=person.id))
    session.commit()
    return profile(person)


@app.get('/events/{event_id}/admins')
def admins(event_id: int, user: User = Depends(superuser), session: Session = Depends(db)):
    get_event(session, event_id)
    return [profile(u) for u in session.scalars(select(User).join(EventAdmin).where(EventAdmin.event_id == event_id))]


@app.delete('/events/{event_id}/admins/{user_id}', status_code=204)
def remove_admin(event_id: int, user_id: int, user: User = Depends(superuser), session: Session = Depends(db)):
    get_event(session, event_id)
    relation = session.get(EventAdmin, (event_id, user_id))
    if not relation:
        raise HTTPException(404, 'Không tìm thấy admin')
    session.delete(relation)
    session.commit()


@app.post('/events/{event_id}/students', status_code=201)
def add_student(event_id: int, body: Person, user: User = Depends(current), session: Session = Depends(db)):
    manage(session, user, event_id)
    person = ensure_user(session, body.email, body.name)
    if session.get(Registration, (event_id, person.id)):
        raise HTTPException(409, 'Sinh viên đã đăng ký sự kiện')
    session.add(Registration(event_id=event_id, user_id=person.id))
    session.commit()
    return profile(person)


def workbook_response(headers, rows, filename):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Students'
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='123B78')
    for row in rows:
        sheet.append(row)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = 's'  # Never interpret user-controlled content as formulas.
                cell.number_format = '@'
    for i, width in enumerate([22, 32, 48, 24, 30, 48], 1):
        sheet.column_dimensions[chr(64 + i)].width = width
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    buffer = BytesIO()
    book.save(buffer)
    return Response(buffer.getvalue(), media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': f'attachment; filename="{filename}"'})


@app.get('/events/{event_id}/students/template')
def template(event_id: int, user: User = Depends(current), session: Session = Depends(db)):
    manage(session, user, event_id)
    return workbook_response(['student_id', 'full_name', 'email'], [], 'students-template.xlsx')


def parse_student_excel(content):
    """Validate the entire workbook before either event or club imports write data."""
    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(413, 'File tối đa 5 MB')
    try:
        archive = zipfile.ZipFile(BytesIO(content))
        if sum(x.file_size for x in archive.infolist()) > 30 * 1024 * 1024:
            raise ValueError('File giải nén quá lớn')
        book = load_workbook(BytesIO(content), read_only=True, data_only=False)
        sheet = book.active
        if sheet.max_row > 10001 or sheet.max_column > 20:
            raise ValueError('Tối đa 10.000 dòng và 20 cột')
        iterator = sheet.iter_rows(values_only=True)
        if tuple(next(iterator, ())) != ('student_id', 'full_name', 'email'):
            raise ValueError('Sai mẫu: cần đúng 3 cột student_id, full_name, email')
        valid, errors, duplicates, seen = [], [], [], set()
        for row_number, row in enumerate(iterator, 2):
            if row_number > 10001:
                raise ValueError('Tối đa 10.000 dòng')
            if all(value is None for value in row):
                continue
            sid, name, email = [str(value).strip() if value is not None else '' for value in row]
            try:
                email = email or f'{sid}@{DOMAIN}'
                email, derived = student_email(email)
                if sid != derived or not name or len(name) > 200:
                    raise ValueError('MSSV/email không khớp hoặc họ tên không hợp lệ; MSSV phải để dạng Text')
                if email in seen:
                    duplicates.append({'row': row_number, 'student_id': sid, 'reason': 'Trùng trong file'})
                    continue
                seen.add(email)
                valid.append((row_number, email, name))
            except (HTTPException, ValueError) as exc:
                errors.append({'row': row_number, 'error': str(exc.detail if isinstance(exc, HTTPException) else exc)})
        book.close()
    except (ValueError, zipfile.BadZipFile, KeyError, StopIteration, TypeError, ParseError, OSError, InvalidFileException) as exc:
        raise HTTPException(422, f'Không đọc được Excel: {exc}')
    return valid, errors, duplicates


@app.post('/events/{event_id}/students/import')
def import_students(event_id: int, file: UploadFile = File(...), user: User = Depends(current), session: Session = Depends(db)):
    manage(session, user, event_id)
    content = file.file.read(5 * 1024 * 1024 + 1)
    valid, errors, duplicates = parse_student_excel(content)
    if errors:
        return {'imported': 0, 'errors': errors, 'duplicates': duplicates, 'message': 'Chưa nhập dòng nào. Sửa các dòng lỗi rồi tải lại.'}
    imported = 0
    for row_number, email, name in valid:
        person = ensure_user(session, email, name)
        if session.get(Registration, (event_id, person.id)):
            duplicates.append({'row': row_number, 'student_id': person.student_id, 'reason': 'Đã có trong sự kiện'})
        else:
            session.add(Registration(event_id=event_id, user_id=person.id))
            imported += 1
    try:
        session.commit()
    except IntegrityError:
        session.rollback()
        raise HTTPException(409, 'Có import đồng thời; hãy tải lại file')
    return {'imported': imported, 'errors': [], 'duplicates': duplicates}


def student_rows(session, event_id, search='', attended=None):
    query = select(User, Attendance).join(Registration, Registration.user_id == User.id).outerjoin(Attendance, (Attendance.user_id == User.id) & (Attendance.event_id == event_id)).where(Registration.event_id == event_id)
    if search:
        query = query.where(User.student_id.contains(search, autoescape=True))
    if attended is not None:
        query = query.where(Attendance.id.is_not(None) if attended else Attendance.id.is_(None))
    return query.order_by(User.student_id)


def attendance_data(person, record):
    return {**profile(person), 'attended': record is not None, 'checked_at': record.checked_at.replace(tzinfo=timezone.utc) if record and record.checked_at.tzinfo is None else record.checked_at if record else None, 'checked_by': record.checked_by if record else None}


@app.get('/events/{event_id}/students')
def students(event_id: int, search: str = Query('', max_length=15), attended: bool | None = None, offset: int = Query(0, ge=0), limit: int = Query(50, ge=1, le=500), user: User = Depends(current), session: Session = Depends(db)):
    manage(session, user, event_id)
    return [attendance_data(p, a) for p, a in session.execute(student_rows(session, event_id, search, attended).offset(offset).limit(limit))]


@app.get('/events/{event_id}/attendance/export')
def export(event_id: int, search: str = Query('', max_length=15), attended: bool | None = None, user: User = Depends(current), session: Session = Depends(db)):
    manage(session, user, event_id)
    rows = []
    for person, record in session.execute(student_rows(session, event_id, search, attended)):
        scanner = session.get(User, record.checked_by) if record else None
        rows.append([person.student_id, person.name, person.email, 'Có mặt' if record else 'Chưa điểm danh', attendance_data(person, record)['checked_at'].isoformat() if record else '', scanner.email if scanner else ''])
    return workbook_response(['student_id', 'full_name', 'email', 'status', 'checked_at_utc', 'checked_by'], rows, f'attendance-{event_id}.xlsx')


@app.get('/events/{event_id}/qr')
def my_qr(event_id: int, user: User = Depends(current), session: Session = Depends(db)):
    item = get_event(session, event_id)
    if not session.get(Registration, (event_id, user.id)):
        raise HTTPException(403, 'Bạn chưa đăng ký sự kiện này')
    if not item.checkin_open:
        raise HTTPException(409, 'Sự kiện chưa mở điểm danh')
    value = token({'sub': str(user.id), 'kind': 'checkin', 'event_id': event_id}, settings.qr_seconds)
    import base64
    image = qrcode.make(value)
    buffer = BytesIO()
    image.save(buffer, format='PNG')
    return {'token': value, 'expires_in': settings.qr_seconds, 'image_data_url': 'data:image/png;base64,' + base64.b64encode(buffer.getvalue()).decode()}


class Scan(BaseModel):
    token: str = Field(min_length=10, max_length=4000)


@app.post('/events/{event_id}/attendance/scan')
def scan(event_id: int, body: Scan, user: User = Depends(current), session: Session = Depends(db)):
    item = manage(session, user, event_id)
    if not item.checkin_open:
        raise HTTPException(409, 'Sự kiện đang đóng điểm danh')
    claims = decode(body.token, 'checkin')
    if claims.get('event_id') != event_id:
        raise HTTPException(403, 'QR thuộc sự kiện khác')
    try:
        sid = int(claims['sub'])
    except (ValueError, TypeError):
        raise HTTPException(401, 'QR không hợp lệ')
    if not session.get(Registration, (event_id, sid)):
        raise HTTPException(403, 'Sinh viên không thuộc sự kiện')
    person = session.get(User, sid)
    record = session.scalar(select(Attendance).where(Attendance.event_id == event_id, Attendance.user_id == sid))
    if record:
        return {'already_checked_in': True, **attendance_data(person, record)}
    record = Attendance(event_id=event_id, user_id=sid, checked_by=user.id)
    session.add(record)
    try:
        session.commit()
    except IntegrityError:
        session.rollback()
        record = session.scalar(select(Attendance).where(Attendance.event_id == event_id, Attendance.user_id == sid))
        if not record:
            raise HTTPException(409, 'Xung đột ghi dữ liệu; hãy quét lại')
        return {'already_checked_in': True, **attendance_data(person, record)}
    return {'already_checked_in': False, **attendance_data(person, record)}


@app.get('/health')
def health():
    return {'status': 'ok'}


@app.post('/superadmins', status_code=201)
def add_superadmin(body: Person, user: User = Depends(superuser), session: Session = Depends(db)):
    person = ensure_user(session, body.email, body.name)
    if is_super(person):
        raise HTTPException(409, 'Tài khoản đã là superadmin')
    session.add(SuperAdmin(user_id=person.id))
    session.commit()
    return profile(person)


@app.get('/superadmins')
def list_superadmins(user: User = Depends(superuser), session: Session = Depends(db)):
    configured = [x.strip().lower() for x in settings.superadmin_emails.split(',') if x.strip()]
    people = session.scalars(select(User).where(or_(User.email.in_(configured), User.id.in_(select(SuperAdmin.user_id)))))
    return {'bootstrap_emails': configured, 'users': [profile(p) for p in people]}


@app.delete('/superadmins/{user_id}', status_code=204)
def revoke_superadmin(user_id: int, user: User = Depends(superuser), session: Session = Depends(db)):
    person = session.get(User, user_id)
    if person and person.email in {x.strip().lower() for x in settings.superadmin_emails.split(',') if x.strip()}:
        raise HTTPException(409, 'Superadmin khởi tạo quản lý qua cấu hình máy chủ')
    role = session.get(SuperAdmin, user_id)
    if not role:
        raise HTTPException(404, 'Không tìm thấy superadmin')
    session.delete(role)
    session.commit()


from .clubs import Club, ClubAdmin, router as club_router
app.include_router(club_router)

from .frontend import register_frontend
register_frontend(app)
